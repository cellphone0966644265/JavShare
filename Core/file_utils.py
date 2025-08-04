# /core/file_utils.py
import os
import pandas as pd
import openpyxl
import argparse
import json

def get_name_in_excel_no_ext(ws, table_name, host):
    """Lấy danh sách tên file (không có phần mở rộng) từ Excel cho một host cụ thể."""
    table = ws.tables[table_name]
    data_range = ws[table.ref]
    rows = list(data_range)
    columns = [cell.value for cell in rows[0]]
    data = [[cell.value for cell in row] for row in rows[1:]]
    df = pd.DataFrame(data, columns=columns)
    
    filtered_df = df[df['Host'] == host]
    return {str(name).rsplit('.', 1)[0] for name in filtered_df['Name']}

def main():
    parser = argparse.ArgumentParser(description="Lấy danh sách file cần upload.")
    parser.add_argument("--source-dir", required=True)
    parser.add_argument("--excel-file", required=True)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--table", required=True)
    parser.add_argument("--host", required=True)
    args = parser.parse_args()

    try:
        wb = openpyxl.load_workbook(args.excel_file, data_only=True)
        ws = wb[args.sheet]
        
        uploaded_files_no_ext = get_name_in_excel_no_ext(ws, args.table, args.host)
        
        files_to_upload = []
        for root, dirs, files in os.walk(args.source_dir):
            for file in files:
                file_name_no_ext = os.path.basename(file).rsplit('.', 1)[0]
                if file_name_no_ext not in uploaded_files_no_ext:
                    files_to_upload.append(os.path.join(root, file))
        
        print(json.dumps(files_to_upload, indent=4))

    except Exception as e:
        print(json.dumps({"status": "error", "message": str(e)}, indent=4))
        exit(1)

if __name__ == '__main__':
    main()
